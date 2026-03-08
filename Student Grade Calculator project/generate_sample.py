#!/usr/bin/env python3
"""
generate_sample.py
──────────────────
Generates a sample input Excel file for GradeCalc Pro testing.
Run: pip install openpyxl && python generate_sample.py
"""

import random
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

STUDENTS = [
    ("Aline Mbarga",       "ETU2024001"),
    ("Boris Nkeng",        "ETU2024002"),
    ("Céline Fotso",       "ETU2024003"),
    ("David Tchamba",      "ETU2024004"),
    ("Estelle Nguefack",   "ETU2024005"),
    ("Franck Kamga",       "ETU2024006"),
    ("Grace Abena",        "ETU2024007"),
    ("Henri Essomba",      "ETU2024008"),
    ("Irène Mfou",         "ETU2024009"),
    ("Jules Ondoa",        "ETU2024010"),
    ("Karine Bello",       "ETU2024011"),
    ("Laurent Ngnié",      "ETU2024012"),
    ("Marie Djomo",        "ETU2024013"),
    ("Narcisse Owono",     "ETU2024014"),
    ("Olivia Menye",       "ETU2024015"),
    ("Patrick Zang",       "ETU2024016"),
    ("Queenie Ateba",      "ETU2024017"),
    ("Roland Evehe",       "ETU2024018"),
    ("Sandra Fouda",       "ETU2024019"),
    ("Thomas Bikié",       "ETU2024020"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Notes"

# Header style
header_fill = PatternFill("solid", fgColor="4F6AF5")
header_font = Font(bold=True, color="FFFFFF", size=12)
center      = Alignment(horizontal="center")

headers = ["Nom", "Matricule", "Note"]
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill      = header_fill
    cell.font      = header_font
    cell.alignment = center

# Data rows with a realistic distribution of grades
random.seed(42)
for i, (name, mat) in enumerate(STUDENTS, 2):
    note = round(random.gauss(62, 18), 1)
    note = max(0, min(100, note))
    ws.cell(row=i, column=1, value=name)
    ws.cell(row=i, column=2, value=mat)
    cell = ws.cell(row=i, column=3, value=round(note, 1))
    cell.alignment = center

# Column widths
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 10

wb.save("sample_notes.xlsx")
print("✅ sample_notes.xlsx créé avec", len(STUDENTS), "étudiants.")
